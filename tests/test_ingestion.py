from __future__ import annotations

import io
import subprocess
from email.message import EmailMessage
from pathlib import Path

import requests
import xlwt
from docx import Document

from shipping_bill_agent.config import Settings
from shipping_bill_agent.demo_data import build_sample_email
from shipping_bill_agent.ingestion.context import ParserContext
from shipping_bill_agent.ingestion.pdf import parse_pdf_bytes
from shipping_bill_agent.ingestion.router import DocumentParser
from shipping_bill_agent.ocr.client import OCRClient, ocr_response_to_blocks


def test_eml_and_excel_attachment_are_parsed(tmp_path) -> None:
    path = tmp_path / "bill.eml"
    path.write_bytes(build_sample_email())
    result = DocumentParser(Settings()).parse_path(path)

    assert len(result.sources) == 2
    attachment = next(source for source in result.sources if source.source_type == "excel")
    assert attachment.blocks
    assert attachment.blocks[0].cells[0] == ["Charge Item", "Amount", "Currency", "Payment Term"]


def test_single_part_eml_body_is_parsed(tmp_path) -> None:
    message = EmailMessage()
    message["Subject"] = "Simple bill"
    message["From"] = "Origin Agent <ops@origin.example>"
    message["To"] = "Destination Agent <billing@destination.example>"
    message.set_content("Document Type: Debit Note\nH/D Fee    40.00 USD PP")
    path = tmp_path / "simple.eml"
    path.write_bytes(message.as_bytes())

    result = DocumentParser(Settings()).parse_path(path)
    body = next(source for source in result.sources if source.source_type == "eml")

    assert "H/D Fee" in body.text


def test_docx_paragraph_and_table_are_parsed(tmp_path) -> None:
    document = Document()
    document.add_paragraph("Document Type: Invoice")
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "Charge"
    table.cell(0, 1).text = "Amount"
    table.cell(0, 2).text = "Currency"
    table.cell(1, 0).text = "Doc Fee"
    table.cell(1, 1).text = "25"
    table.cell(1, 2).text = "USD"
    buffer = io.BytesIO()
    document.save(buffer)
    path = tmp_path / "bill.docx"
    path.write_bytes(buffer.getvalue())

    result = DocumentParser(Settings()).parse_path(path)
    assert len(result.sources) == 1
    assert any(block.cells for block in result.sources[0].blocks)


def test_legacy_doc_uses_conversion_path(monkeypatch, tmp_path) -> None:
    from shipping_bill_agent.ingestion import word

    document = Document()
    document.add_paragraph("Document Type: Invoice")
    buffer = io.BytesIO()
    document.save(buffer)
    converted_docx = buffer.getvalue()

    monkeypatch.setattr(word, "_convert_legacy_doc", lambda _data: converted_docx)
    path = tmp_path / "legacy.doc"
    path.write_bytes(b"legacy-word-document")

    result = DocumentParser(Settings()).parse_path(path)

    assert result.sources[0].source_type == "word"
    assert result.sources[0].blocks[0].text == "Document Type: Invoice"


def test_word_com_conversion_command(monkeypatch) -> None:
    from shipping_bill_agent.ingestion import word

    converted = b"converted-docx"

    def fake_run(command, **kwargs):
        assert command[0] == "powershell"
        assert "-NoProfile" in command
        assert "Word.Application" in command[-1]
        Path(kwargs["env"]["SHIPPING_WORD_DESTINATION"]).write_bytes(converted)
        return subprocess.CompletedProcess(command, 0, "", "")

    monkeypatch.setattr(word.subprocess, "run", fake_run)

    assert word._convert_with_word(b"legacy-doc", "powershell") == converted


def test_legacy_doc_prefers_word_then_falls_back(monkeypatch) -> None:
    from shipping_bill_agent.ingestion import word
    from shipping_bill_agent.ingestion.exceptions import DocumentParseError

    def fail_word(_data, _executable):
        raise DocumentParseError("Word COM 转换失败")

    monkeypatch.setattr(word, "_word_powershell", lambda: "powershell")
    monkeypatch.setattr(word, "_convert_with_word", fail_word)
    monkeypatch.setattr(word, "_convert_with_libreoffice", lambda _data: b"libreoffice-docx")

    assert word._convert_legacy_doc(b"legacy-doc") == b"libreoffice-docx"


def test_legacy_xls_is_parsed(tmp_path) -> None:
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Charges")
    for column, value in enumerate(["Charge", "Amount", "Currency"]):
        sheet.write(0, column, value)
    for column, value in enumerate(["THC", 120, "USD"]):
        sheet.write(1, column, value)
    path = tmp_path / "bill.xls"
    workbook.save(str(path))

    result = DocumentParser(Settings()).parse_path(path)
    assert result.sources[0].blocks[0].cells[1] == ["THC", "120.0", "USD"]


def test_nested_eml_attachment_is_parsed(tmp_path) -> None:
    from shipping_bill_agent.demo_data import build_charge_workbook

    inner = EmailMessage()
    inner["Subject"] = "Nested bill"
    inner["From"] = "Origin Agent <ops@origin.example>"
    inner["To"] = "Destination Agent <billing@destination.example>"
    inner.set_content("Attached bill")
    inner.add_attachment(
        build_charge_workbook(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="nested.xlsx",
    )

    outer = EmailMessage()
    outer["Subject"] = "Forwarded bill"
    outer["From"] = "Ops <ops@origin.example>"
    outer["To"] = "Billing <billing@destination.example>"
    outer.set_content("Forwarding the bill")
    outer.add_attachment(inner)

    path = tmp_path / "outer.eml"
    path.write_bytes(outer.as_bytes())
    result = DocumentParser(Settings()).parse_path(path)
    assert {source.source_type for source in result.sources} == {"eml", "excel"}
    assert len(result.sources) == 3


def test_ocr_response_preserves_table_cells() -> None:
    blocks = ocr_response_to_blocks(
        {"result": {"tables": [[["Charge", "Amount"], ["THC", "120.00"]]]}},
        "test",
    )
    assert len(blocks) == 1
    assert blocks[0].cells[1] == ["THC", "120.00"]


def test_empty_ocr_response_does_not_become_fake_text() -> None:
    blocks = ocr_response_to_blocks(
        {
            "data": {
                "createDate": "2026-10-20 22:03:05",
                "fileName": "",
                "list": [],
                "orderType": "None",
            },
            "success": True,
            "message": "文档处理成功",
        },
        "test",
    )
    assert blocks == []


def test_ocr_response_preserves_text_and_table_confidence() -> None:
    blocks = ocr_response_to_blocks(
        {
            "data": {
                "texts": [
                    {"text": "THC 120.00 USD", "score": 95, "pageNo": 2},
                    {"text": "Tax 5.00 USD", "confidence": 0.62, "page": 3},
                ],
                "tables": [
                    {
                        "cells": [["Charge", "Amount"], ["THC", "120.00"]],
                        "confidence": 0.91,
                        "page_number": 2,
                    }
                ],
            }
        },
        "test",
    )

    assert [block.confidence for block in blocks] == [0.95, 0.62, 0.91]
    assert [block.page_no for block in blocks] == [2, 3, 2]


def test_pdf_empty_ocr_response_is_recorded_as_warning() -> None:
    class EmptyOCR:
        def recognize_bytes(self, data, file_name, mime_type):
            return []

    path = Path("笔试题AIAGENT") / "PDF 示例账单2.pdf"
    source = parse_pdf_bytes(
        path.read_bytes(),
        path.name,
        ParserContext(settings=Settings(), ocr_client=EmptyOCR()),
    )[0]

    assert "OCR 未返回可识别内容" in source.metadata["warnings"]


def test_ocr_client_uses_required_request_shape() -> None:
    class Publisher:
        def publish(self, data, file_name, mime_type):
            assert data == b"image"
            assert file_name == "bill.png"
            return "https://files.example/bill.png"

    class Response:
        def raise_for_status(self):
            return None

        def json(self):
            return {"text": "THC 120.00 USD"}

    class Session:
        def __init__(self):
            self.request = None

        def post(self, endpoint, json, timeout):
            self.request = (endpoint, json, timeout)
            return Response()

    session = Session()
    client = OCRClient(
        "http://ocr.example/process",
        Publisher(),
        task="orderChange",
        session=session,
    )
    blocks = client.recognize_bytes(b"image", "bill.png", "image/png")
    assert blocks[0].text == "THC 120.00 USD"
    assert session.request[1] == {
        "file": "https://files.example/bill.png",
        "task": "orderChange",
        "callBack": "",
        "taskType": "",
    }


def test_ocr_client_retries_transport_failure() -> None:
    class Publisher:
        def publish(self, data, file_name, mime_type):
            return "https://files.example/bill.png"

    class Response:
        def raise_for_status(self):
            return None

        def json(self):
            return {"text": "Storage Fee 10.00 USD"}

    class FlakySession:
        def __init__(self):
            self.calls = 0

        def post(self, endpoint, json, timeout):
            self.calls += 1
            if self.calls == 1:
                raise requests.ConnectionError("temporary failure")
            return Response()

    session = FlakySession()
    client = OCRClient(
        "http://ocr.example/process",
        Publisher(),
        session=session,
        max_retries=1,
    )
    blocks = client.recognize_bytes(b"image", "bill.png", "image/png")
    assert session.calls == 2
    assert blocks[0].text == "Storage Fee 10.00 USD"
