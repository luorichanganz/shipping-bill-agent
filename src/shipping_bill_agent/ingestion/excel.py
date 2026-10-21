from __future__ import annotations

import hashlib
import io

import xlrd
from openpyxl import load_workbook

from shipping_bill_agent.ingestion.context import ParserContext
from shipping_bill_agent.ingestion.exceptions import DocumentParseError
from shipping_bill_agent.models.document import BlockType, ContentBlock, SourceDocument


def parse_excel_bytes(data: bytes, file_name: str, context: ParserContext) -> list[SourceDocument]:
    digest = hashlib.sha256(data).hexdigest()
    source_id = f"excel-{digest[:16]}"
    if file_name.lower().endswith(".xls"):
        return _parse_legacy_excel(data, file_name, source_id, digest)
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    except Exception as exc:
        raise DocumentParseError(f"无法读取 Excel 文件: {exc}") from exc

    blocks: list[ContentBlock] = []
    sheet_names: list[str] = []
    for sheet in workbook.worksheets:
        sheet_names.append(sheet.title)
        rows: list[list[str | None]] = []
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() if value is not None else None for value in row]
            if any(values):
                rows.append(values)
        if rows:
            blocks.append(
                ContentBlock(
                    block_id=f"{source_id}:{sheet.title}:table",
                    block_type=BlockType.TABLE,
                    sheet_name=sheet.title,
                    row_no=1,
                    cells=rows,
                )
            )
    workbook.close()
    return [
        SourceDocument(
            source_id=source_id,
            source_type="excel",
            file_name=file_name,
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            sha256=digest,
            parent_source_id=None,
            metadata={"sheets": sheet_names},
            blocks=blocks,
        )
    ]


def _parse_legacy_excel(
    data: bytes,
    file_name: str,
    source_id: str,
    digest: str,
) -> list[SourceDocument]:
    try:
        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
    except Exception as exc:
        raise DocumentParseError(f"无法读取 XLS 文件: {exc}") from exc

    blocks: list[ContentBlock] = []
    sheet_names = workbook.sheet_names()
    for sheet_name in sheet_names:
        sheet = workbook.sheet_by_name(sheet_name)
        rows: list[list[str | None]] = []
        for row_index in range(sheet.nrows):
            values = []
            for cell in sheet.row(row_index):
                if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                    values.append(None)
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    values.append(xlrd.xldate_as_datetime(cell.value, workbook.datemode).date().isoformat())
                else:
                    values.append(str(cell.value).strip())
            if any(values):
                rows.append(values)
        if rows:
            blocks.append(
                ContentBlock(
                    block_id=f"{source_id}:{sheet_name}:table",
                    block_type=BlockType.TABLE,
                    sheet_name=sheet_name,
                    row_no=1,
                    cells=rows,
                )
            )
    workbook.release_resources()
    return [
        SourceDocument(
            source_id=source_id,
            source_type="excel",
            file_name=file_name,
            mime_type="application/vnd.ms-excel",
            sha256=digest,
            parent_source_id=None,
            metadata={"sheets": sheet_names},
            blocks=blocks,
        )
    ]
