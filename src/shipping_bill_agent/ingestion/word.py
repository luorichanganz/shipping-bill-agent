from __future__ import annotations

import hashlib
import io
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document

from shipping_bill_agent.ingestion.context import ParserContext
from shipping_bill_agent.ingestion.exceptions import DocumentParseError
from shipping_bill_agent.models.document import BlockType, ContentBlock, SourceDocument

_WORD_COM_SCRIPT = r"""
$ErrorActionPreference = 'Stop'
$word = $null
$document = $null
try {
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $word.DisplayAlerts = 0
    try { $word.AutomationSecurity = 3 } catch {}
    $document = $word.Documents.Open(
        $env:SHIPPING_WORD_SOURCE,
        $false,
        $true,
        $false
    )
    $document.SaveAs2($env:SHIPPING_WORD_DESTINATION, 16)
}
finally {
    if ($null -ne $document) { $document.Close(0) }
    if ($null -ne $word) { $word.Quit() }
    if ($null -ne $document) {
        [Runtime.InteropServices.Marshal]::ReleaseComObject($document) | Out-Null
    }
    if ($null -ne $word) {
        [Runtime.InteropServices.Marshal]::ReleaseComObject($word) | Out-Null
    }
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
}
"""


def _parse_docx_data(data: bytes, source_id: str) -> list[ContentBlock]:
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentParseError(f"无法读取 Word 文件: {exc}") from exc

    blocks: list[ContentBlock] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if text:
            blocks.append(
                ContentBlock(
                    block_id=f"{source_id}:paragraph{index}",
                    block_type=BlockType.TEXT,
                    text=text,
                    row_no=index,
                )
            )
    for table_index, table in enumerate(document.tables, start=1):
        cells = [[cell.text.strip() or None for cell in row.cells] for row in table.rows]
        blocks.append(
            ContentBlock(
                block_id=f"{source_id}:table{table_index}",
                block_type=BlockType.TABLE,
                cells=cells,
            )
        )
    return blocks


def _word_powershell() -> str | None:
    if os.name != "nt":
        return None
    return shutil.which("powershell") or shutil.which("pwsh")


def _convert_with_word(data: bytes, executable: str) -> bytes:
    with tempfile.TemporaryDirectory(prefix="shipping-word-") as temp_dir:
        source = Path(temp_dir) / "source.doc"
        destination = Path(temp_dir) / "source.docx"
        source.write_bytes(data)
        environment = os.environ.copy()
        environment["SHIPPING_WORD_SOURCE"] = str(source)
        environment["SHIPPING_WORD_DESTINATION"] = str(destination)
        completed = subprocess.run(
            [
                executable,
                "-NoProfile",
                "-NonInteractive",
                "-Command",
                _WORD_COM_SCRIPT,
            ],
            check=False,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=60,
            env=environment,
        )
        if completed.returncode != 0:
            detail = (completed.stderr or completed.stdout).strip()
            raise DocumentParseError(f"Word COM 转换失败: {detail or '未知错误'}")
        if not destination.exists() or destination.stat().st_size == 0:
            raise DocumentParseError("Word COM 未生成 .docx 文件")
        return destination.read_bytes()


def _convert_with_libreoffice(data: bytes) -> bytes:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        raise DocumentParseError("旧版 .doc 需要 LibreOffice 的 soffice/libreoffice 进行转换；当前环境未安装")
    with tempfile.TemporaryDirectory(prefix="shipping-doc-") as temp_dir:
        source = Path(temp_dir) / "source.doc"
        source.write_bytes(data)
        subprocess.run(
            [
                executable,
                "--headless",
                "--convert-to",
                "docx",
                "--outdir",
                temp_dir,
                str(source),
            ],
            check=True,
            capture_output=True,
            timeout=60,
        )
        converted = Path(temp_dir) / "source.docx"
        if not converted.exists():
            raise DocumentParseError("LibreOffice 未能将 .doc 转换为 .docx")
        return converted.read_bytes()


def _convert_legacy_doc(data: bytes) -> bytes:
    errors: list[str] = []
    word_executable = _word_powershell()
    if word_executable:
        try:
            return _convert_with_word(data, word_executable)
        except Exception as exc:
            errors.append(str(exc))
    try:
        return _convert_with_libreoffice(data)
    except Exception as exc:
        errors.append(str(exc))
    details = "；".join(errors)
    raise DocumentParseError(f"旧版 .doc 自动转换失败: {details}")


def parse_word_bytes(data: bytes, file_name: str, context: ParserContext) -> list[SourceDocument]:
    digest = hashlib.sha256(data).hexdigest()
    source_id = f"word-{digest[:16]}"
    if file_name.lower().endswith(".doc"):
        data = _convert_legacy_doc(data)
    blocks = _parse_docx_data(data, source_id)
    return [
        SourceDocument(
            source_id=source_id,
            source_type="word",
            file_name=file_name,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            sha256=digest,
            parent_source_id=None,
            metadata={},
            blocks=blocks,
        )
    ]
